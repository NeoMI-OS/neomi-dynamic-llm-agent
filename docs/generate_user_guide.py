"""
NEOMI Pipeline Runner — Felhasználói Kézikönyv generátor
Futtasd a repo gyökeréből: python3 docs/generate_user_guide.py
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

SCREENSHOTS = os.path.join(os.path.dirname(__file__), '..', 'tests', 'ui', 'screenshots')
OUTPUT = os.path.join(os.path.dirname(__file__), 'NEOMI_Felhasznaloi_Kezikonyv.docx')

def add_screenshot(doc, filename, caption):
    path = os.path.join(SCREENSHOTS, filename)
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(5.5))
        p = doc.add_paragraph(caption)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        p.runs[0].font.italic = True
    else:
        p = doc.add_paragraph(f'[Hiányzó kép: {filename}]')
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0x44, 0x44)

def heading1(doc, text):
    h = doc.add_heading(text, level=1)
    h.runs[0].font.color.rgb = RGBColor(0x00, 0xCC, 0x66)

def heading2(doc, text):
    h = doc.add_heading(text, level=2)
    h.runs[0].font.color.rgb = RGBColor(0x33, 0x99, 0xFF)

def heading3(doc, text):
    doc.add_heading(text, level=3)

def body(doc, text):
    doc.add_paragraph(text)

def note(doc, text):
    p = doc.add_paragraph()
    p.add_run('ℹ️  ').bold = True
    r = p.add_run(text)
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

def table_row(t, cells, bold=False):
    row = t.add_row()
    for i, text in enumerate(cells):
        cell = row.cells[i]
        cell.text = text
        if bold:
            for run in cell.paragraphs[0].runs:
                run.bold = True

def build():
    doc = Document()

    # Margók
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # Cím
    title = doc.add_heading('NEOMI Pipeline Runner', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph('Felhasználói Kézikönyv')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    meta = doc.add_paragraph('Verzió: 1.0  |  2026. május  |  NeoMI projekt')
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.runs[0].font.size = Pt(10)
    meta.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_page_break()

    # ── 1. Áttekintés ──────────────────────────────────────────────────────────
    heading1(doc, '1. Áttekintés')
    body(doc,
        'A NEOMI Pipeline Runner egy mesterséges intelligencia alapú tartalomgeneráló eszköz, '
        'amellyel vállalati AI-oktatási anyagokat lehet automatikusan létrehozni. '
        'A rendszer egy 5 lépéses ügynök-pipeline-t futtat:\n\n'
        '  1. Kontextuselemzés\n'
        '  2. Igényfelmérés\n'
        '  3. Tantervtervezés\n'
        '  4. Tartalomírás\n'
        '  5. Kritikai értékelés'
    )
    body(doc, 'Az alkalmazás elérhető a következő URL-en:')
    p = doc.add_paragraph('https://neomi-pipeline-ui-643234238822.europe-west1.run.app')
    p.runs[0].font.color.rgb = RGBColor(0x00, 0x88, 0xFF)

    # ── 2. Felület felépítése ──────────────────────────────────────────────────
    heading1(doc, '2. A felhasználói felület felépítése')
    body(doc,
        'Az alkalmazás két fő területre osztott: a bal oldalon találhatók a bemeneti adatok, '
        'a jobb oldalon jelennek meg a pipeline eredményei.'
    )

    add_screenshot(doc, '01_initial_load.png', '1. ábra — Kezdőképernyő, üres állapot')
    doc.add_paragraph()

    body(doc, 'A fejlécben látható elemek:')
    items = doc.add_paragraph()
    items.style = 'List Bullet'
    items.text = 'Zöld villogó pont — a rendszer aktív állapotát jelzi'
    doc.add_paragraph('NEOMI Pipeline Runner — az alkalmazás neve', style='List Bullet')
    doc.add_paragraph('// exp-001 — a kiválasztott kísérlet azonosítója (jobb felső sarok)', style='List Bullet')

    # ── 3. Bal panel ──────────────────────────────────────────────────────────
    heading1(doc, '3. Bal panel — Bemeneti adatok')

    heading2(doc, '3.1 Kísérlet kiválasztása')
    body(doc,
        'A KÍSÉRLET szekcióban választható ki a pipeline LLM-konfigurációja. '
        'Három előre definiált kísérlet érhető el:'
    )

    t = doc.add_table(rows=1, cols=2)
    t.style = 'Table Grid'
    table_row(t, ['Kísérlet', 'Leírás'], bold=True)
    table_row(t, ['exp-001 — Baseline, GPT-4o mindenhol', 'Referencia pont: azonos erős modell minden szerepben'])
    table_row(t, ['exp-002 — Claude diverzitás', 'Claude Sonnet minden pipeline csomópontban'])
    table_row(t, ['exp-003 — Vegyes optimális', 'Különböző modellek az egyes szerepekhez optimalizálva'])
    doc.add_paragraph()

    heading2(doc, '3.2 Bemeneti mezők')
    add_screenshot(doc, '02_experiment_selector.png', '2. ábra — Bemeneti form, kötelező mező figyelmeztetéssel')
    doc.add_paragraph()

    body(doc, 'Az alábbi mezőket lehet kitölteni:')
    t2 = doc.add_table(rows=1, cols=3)
    t2.style = 'Table Grid'
    table_row(t2, ['Mező', 'Kötelező', 'Leírás'], bold=True)
    table_row(t2, ['CÉG NEVE', '✅ Igen', 'A vállalat neve (pl. Acme Solutions Kft.)'])
    table_row(t2, ['CÉG TÍPUSA', 'Nem', 'Az iparág vagy tevékenység típusa'])
    table_row(t2, ['CÉLKÖZÖNSÉG', 'Nem', 'Kiknek szól a tartalom (pl. Középvezetők)'])
    table_row(t2, ['CÉL', 'Nem', 'Mi a képzés célja'])
    table_row(t2, ['KONTEXTUS', 'Nem', 'További háttérinformáció (opcionális)'])
    doc.add_paragraph()

    note(doc,
        'A Cég neve mező kitöltése kötelező. Amíg üres, a „PIPELINE FUTTATÁSA" gomb inaktív '
        '(szürke), és „A cég neve kötelező" figyelmeztetés látható.'
    )

    heading2(doc, '3.3 Kitöltött állapot')
    add_screenshot(doc, '03_company_name_filled.png', '3. ábra — Cégnév kitöltve, a Run gomb aktívvá válik')
    doc.add_paragraph()
    add_screenshot(doc, '04_all_fields_filled.png', '4. ábra — Összes mező kitöltve, pipeline indítható')
    doc.add_paragraph()

    body(doc,
        'Ha a Cég neve mező ki van töltve, a „PIPELINE FUTTATÁSA" gomb zöldre vált és kattinthatóvá válik.'
    )

    # ── 4. Pipeline futtatása ──────────────────────────────────────────────────
    heading1(doc, '4. A pipeline futtatása')

    heading2(doc, '4.1 Indítás')
    body(doc,
        'Kattints a zöld „PIPELINE FUTTATÁSA" gombra. A gomb felirata azonnal '
        '„PIPELINE FUT..."-ra változik és letiltódik, hogy megelőzze a dupla küldést.'
    )

    heading2(doc, '4.2 Betöltési állapot')
    add_screenshot(doc, '05_pipeline_loading.png', '5. ábra — Pipeline futás közbeni betöltési állapot')
    doc.add_paragraph()

    body(doc, 'A jobb panelen megjelenik:')
    doc.add_paragraph('Animált zöld pontok — a feldolgozás folyamatban van', style='List Bullet')
    doc.add_paragraph('„Pipeline fut..." felirat', style='List Bullet')
    doc.add_paragraph('Tájékoztató szöveg a várható futási időről', style='List Bullet')

    note(doc,
        'A pipeline futtatása általában 30–90 másodpercig tart, mivel 5 LLM ügynök dolgozik '
        'sorban egymás után. Ne zárd be a böngésző lapot!'
    )

    # ── 5. Eredmények ──────────────────────────────────────────────────────────
    heading1(doc, '5. Az eredmények értelmezése')

    heading2(doc, '5.1 Pipeline nyomkövetés (Node Trace)')
    body(doc,
        'A pipeline lefutása után a jobb panelen megjelennek az eredmények. '
        'Az összesítő sorban látható a csomópontok száma, a teljes futási idő és az összes felhasznált token. '
        'Alatta minden egyes ügynök egy-egy kártyán jelenik meg.'
    )
    body(doc, 'Egy kártya elemei:')
    doc.add_paragraph('Színes pont — zöld: sikeres, sárga: közepes, piros: hiba', style='List Bullet')
    doc.add_paragraph('Csomópont neve — az ügynök szerepe (pl. context_analyst)', style='List Bullet')
    doc.add_paragraph('Model — a használt LLM modell neve', style='List Bullet')
    doc.add_paragraph('Idő — az ügynök futási ideje másodpercben', style='List Bullet')
    doc.add_paragraph('Token — a felhasznált tokenek száma', style='List Bullet')
    doc.add_paragraph('SRS badge — Service Readiness Score (0–100)', style='List Bullet')
    doc.add_paragraph()
    body(doc, 'Kattints bármelyik kártyára a részletes LLM kimenet megtekintéséhez.')

    heading2(doc, '5.2 SRS (Service Readiness Score)')
    body(doc,
        'Az SRS egy összetett minőségmutató, amely az Artificial Analysis benchmark adatain alapul. '
        'A badge fölé húzva az egeret egy tooltip mutatja a részletes bontást.'
    )
    t3 = doc.add_table(rows=1, cols=3)
    t3.style = 'Table Grid'
    table_row(t3, ['Kategória', 'Súly', 'Leírás'], bold=True)
    table_row(t3, ['Tudás', '25%', 'Factual knowledge benchmark (MMLU Pro, GPQA)'])
    table_row(t3, ['Érvelés', '35%', 'Reasoning benchmark (MATH-500, AIME, LiveCodeBench)'])
    table_row(t3, ['Használhatóság', '40%', 'AI Intelligence Index alapú minőségmutató'])
    doc.add_paragraph()

    body(doc, 'Értelmezés: 🟢 70+ Kiváló  |  🟡 40–70 Elfogadható  |  🔴 0–40 Gyenge')

    heading2(doc, '5.3 Kritikai összefoglaló')
    body(doc,
        'A pipeline utolsó eleme a Critic ügynök értékelése. '
        'Megmutatja az érettségi pontszámot (0–100), a kategóriánkénti értékelést '
        '(relevancia, mélység, alkalmazhatóság) és egy szöveges visszajelzést.'
    )

    # ── 6. Hibaelhárítás ──────────────────────────────────────────────────────
    heading1(doc, '6. Hibaelhárítás')
    t4 = doc.add_table(rows=1, cols=3)
    t4.style = 'Table Grid'
    table_row(t4, ['Hiba', 'Lehetséges ok', 'Megoldás'], bold=True)
    table_row(t4, ['"A cég neve kötelező"', 'CÉG NEVE mező üres', 'Töltsd ki a Cég neve mezőt'])
    table_row(t4, ['"Pipeline futtatása sikertelen"', 'Backend nem elérhető', 'Ellenőrizd az internetkapcsolatot'])
    table_row(t4, ['Üres jobb panel', 'Timeout vagy hálózati hiba', 'Töltsd újra az oldalt és futtasd újra'])
    table_row(t4, ['Hosszú betöltés (90s+)', 'Sok token, komplex bemenet', 'Normális jelenség, várj türelmesen'])

    # ── 7. Műszaki info ────────────────────────────────────────────────────────
    heading1(doc, '7. Műszaki információk')
    t5 = doc.add_table(rows=1, cols=2)
    t5.style = 'Table Grid'
    table_row(t5, ['Összetevő', 'Részletek'], bold=True)
    table_row(t5, ['Backend', 'FastAPI + LangGraph, Google Cloud Run (europe-west1)'])
    table_row(t5, ['Frontend', 'React + Vite, Google Cloud Run (europe-west1)'])
    table_row(t5, ['LLM modellek', 'GPT-4o, GPT-4o-mini, Claude Sonnet, Gemini Flash'])
    table_row(t5, ['GitHub', 'github.com/zsigacsaba/neomi-dynamic-llm-agent'])
    table_row(t5, ['Branch', 'feature/multi-agent-pipeline'])

    doc.save(OUTPUT)
    print(f'✅ Dokumentum létrehozva: {OUTPUT}')

if __name__ == '__main__':
    build()
